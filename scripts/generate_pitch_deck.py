"""
AquaAI — Pitch Deck Generator for IDEA MINERS 2026
Builds high-impact, professionally styled presentation matching the official 9-slide template.
"""

import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_pitch_deck():
    template_path = os.path.abspath("docs/Idea-Miners-pitch-template.pptx")
    output_pptx = os.path.abspath("docs/AquaAI_Idea_Miners_Pitch.pptx")
    output_pdf = os.path.abspath("docs/AquaAI_Idea_Miners_Pitch.pdf")
    
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found: {template_path}")
        
    prs = pptx.Presentation(template_path)
    print(f"Loaded template with {len(prs.slides)} slides.")
    
    # Color Palette (AgriTech Emerald & Slate Dark)
    COLOR_PRIMARY = RGBColor(15, 118, 110)    # Deep Emerald (#0F766E)
    COLOR_ACCENT = RGBColor(5, 150, 105)      # Emerald (#059669)
    COLOR_DARK = RGBColor(15, 23, 42)         # Slate 900 (#0F172A)
    COLOR_MUTED = RGBColor(71, 85, 105)       # Slate 600 (#475569)
    COLOR_LIGHT_BG = RGBColor(248, 250, 252)  # Slate 50 (#F8FAFC)
    COLOR_CARD_BG = RGBColor(255, 255, 255)   # White
    COLOR_CARD_BORDER = RGBColor(203, 213, 225) # Slate 300 (#CBD5E1)
    COLOR_HIGHLIGHT = RGBColor(236, 253, 245) # Mint Highlight (#ECFDF5)
    COLOR_WHITE = RGBColor(255, 255, 255)

    def format_title(shape, text):
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Segoe UI"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY

    # =========================================================================
    # SLIDE 1: Название и одна фраза
    # =========================================================================
    slide1 = prs.slides[0]
    format_title(slide1.shapes[0], "1. Название и одна фраза")
    
    tf1 = slide1.shapes[1].text_frame
    tf1.clear()
    
    p = tf1.paragraphs[0]
    p.text = "AquaAI (АкваАй)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(12)
    
    items1 = [
        ("Суть стартапа (One-Liner):", "Цифровая платформа точного управления орошением для фермеров Центральной Азии на основе международного гидрологического стандарта FAO-56 и объяснимого ИИ (Gemini)."),
        ("Миссия:", "Превратить дефицит поливной воды в управляемый экономический актив фермера, снизить потери воды на 20–35% и защитить экосистему региона."),
        ("География и локация:", "Казахстан (г. Алматы / Туркестанская, Кызылординская, Жамбылская и Алматинская области)."),
        ("Стадия и доступность:", "Полнофункциональный, стабильный MVP готов к тестированию в браузере (на русском и казахском языках)."),
        ("Контакты и демо:", "GitHub: github.com/SalaAkh/AquaAI  |  Demo Live: aqua-ai-demo.loca.lt  |  info@aquaai.kz")
    ]
    
    for title, desc in items1:
        p_item = tf1.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(8)

    # =========================================================================
    # SLIDE 2: Проблема
    # =========================================================================
    slide2 = prs.slides[1]
    format_title(slide2.shapes[0], "2. Проблема: Водный дефицит и неконтролируемые потери")
    
    tf2 = slide2.shapes[1].text_frame
    tf2.clear()
    
    items2 = [
        ("Кто страдает:", "Более 40 000 малых и средних крестьянских хозяйств (20–500 га) орошаемого юга Казахстана (кукуруза, хлопок, рис, овощи, сады) в условиях прогрессирующего маловодья рек Сырдарья и Или."),
        ("В чём проблема:", "80–85% забора пресной воды в регионе уходит на сельское хозяйство. До 40% воды теряется впустую из-за затопления полей «на глаз», фильтрации в земляных арыках и полуденного испарения в зной."),
        ("Как решают сейчас:", "Поливают вслепую по графику подачи воды из канала или интуиции поливальщика, без знания научно обоснованной суточной потребности культур в воде (ETc)."),
        ("Сколько это стоит фермеру:", "Переплата от 1.5 до 3 млн ₸ за сезон за электроэнергию/дизель для холостой работы насосов, а также риски потери до 25–30% урожая из-за водного стресса или заболачивания/засоления."),
        ("Почему не подходят аналоги:", "Зарубежные системы (CropX, FieldNET) стоят от $3,000, требуют обязательной покупки дорогих импортных датчиков, не поддерживают казахский язык и не адаптированы к бассейнам Центральной Азии.")
    ]
    
    for title, desc in items2:
        p_item = tf2.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(8)

    # =========================================================================
    # SLIDE 3: Решение
    # =========================================================================
    slide3 = prs.slides[2]
    format_title(slide3.shapes[0], "3. Решение: Интеллектуальный агрономический расчёт")
    
    tf3 = slide3.shapes[1].text_frame
    tf3.clear()
    
    items3 = [
        ("Что делает продукт:", "AquaAI рассчитывает точную суточную потребность поля в поливе за 30 секунд по координатам, культуре и типу почвы без обязательной покупки оборудования на старте."),
        ("Математическое ядро FAO-56:", "Расчёт эталонной и культурной эвапотранспирации (ETc = ET0 × Kc), эффективных осадков и КПД оросительной сети по стандарту Продовольственной и сельскохозяйственной организации ООН."),
        ("Оптимальные окна полива:", "Рекомендация запуска полива в прохладные утренние (05:00–09:00) и ночные (19:00–22:00) часы, что снижает непродуктивное испарение воды на 20–25%."),
        ("Объяснимый ИИ (Google Gemini):", "Разъяснение расчётов на казахском, русском и английском языках с опорой строго на гидрологические цифры (без галлюцинаций и выдуманных данных)."),
        ("Учёт фактического водопотребления:", "Модуль Water Usage Tracking с 3 режимами (по счётчику, по производительности насоса м³/ч × часы, по талону канала) и сравнением «план vs факт».")
    ]
    
    for title, desc in items3:
        p_item = tf3.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(8)

    # =========================================================================
    # SLIDE 4: Рынок
    # =========================================================================
    slide4 = prs.slides[3]
    format_title(slide4.shapes[0], "4. Рынок: Потенциал орошаемых земель Центральной Азии")
    
    tf4 = slide4.shapes[1].text_frame
    tf4.clear()
    
    items4 = [
        ("TAM (Общий объём целевого рынка):", "8.5 млн га орошаемых земель в Центральной Азии (Казахстан, Узбекистан, Кыргызстан, Таджикистан). При тарифе $1.50/га потенциальный объём рынка специализированного софта водоучёта составляет $12.75M в год (с учётом датчиков и консалтинга — свыше $100M)."),
        ("SAM (Доступный рынок Казахстана):", "1.8 млн га орошаемых площадей в Республике Казахстан (~$2.7M ежегодного софтверного рынка подписки)."),
        ("SOM (Достижимый рынок на первые 2 года):", "50 000 га в Туркестанской, Жамбылской и Кызылординской областях. Ожидаемая выручка (ARR): $75 000 (~37.5 млн ₸)."),
        ("Сколько платит один клиент:", "В среднем $1.50 / га (~750 ₸ / га) за поливной сезон. Среднее фермерское хозяйство площадью 100 га платит всего 75 000 ₸ за сезон."),
        ("Расчёт объёма рынка (SOM):", "50 000 га × $1.50 / га в сезон = $75 000 (~37.5 млн тенге) ежегодной выручки.")
    ]
    
    for title, desc in items4:
        p_item = tf4.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(8)

    # =========================================================================
    # SLIDE 5: Бизнес-модель
    # =========================================================================
    slide5 = prs.slides[4]
    format_title(slide5.shapes[0], "5. Бизнес-модель: Сезонная B2B SaaS-подписка")
    
    tf5 = slide5.shapes[1].text_frame
    tf5.clear()
    
    items5 = [
        ("Кто платит:", "Крестьянские и фермерские хозяйства (КХ/ФХ), агрохолдинги, сельскохозяйственные производственные кооперативы (СПК) и ассоциации водопользователей (АВП)."),
        ("Тариф 1 — Freemium (0 ₸):", "До 2 полей (до 50 га) — базовый расчёт FAO-56 на 7 дней, прогноз Open-Meteo, журнал расхода воды. Инструмент массовой виральной лидогенерации."),
        ("Тариф 2 — Про ($1.50 / га в сезон, ~750 ₸/га):", "Неограниченное число полей и насосов, детальная план-факт аналитика отклонений, ИИ-агроном Gemini на казахском и русском, экспорт официальных PDF/CSV отчётов, калькулятор энергозатрат."),
        ("Тариф 3 — Корпоративный ($500–$1,500 в год):", "Для крупных хозяйств и кооперативов: многопользовательский доступ (агроном, бригадир, бухгалтер), выгрузка для госсубсидий и интеграция с расходомерами."),
        ("Регулярность платежей:", "Ежегодная сезонная подписка с возобновлением весной перед началом активного вегетационного полива (апрель-май). Окупаемость для фермера наступает за счёт предотвращения 2–3 лишних включений насоса.")
    ]
    
    for title, desc in items5:
        p_item = tf5.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(8)

    # =========================================================================
    # SLIDE 6: Трекшн
    # =========================================================================
    slide6 = prs.slides[5]
    format_title(slide6.shapes[0], "6. Трекшн: Готовый MVP и инженерная верификация")
    
    tf6 = slide6.shapes[1].text_frame
    tf6.clear()
    
    items6 = [
        ("Что сделано (Готовый рабочий MVP):", "Разработана полнофункциональная платформа на Next.js 16, React 19, TypeScript, интегрированная с API Open-Meteo, расчётным движком FAO-56 и Google Gemini AI."),
        ("Надёжная база данных и безопасность:", "Облачная СУБД Supabase PostgreSQL с Row Level Security (RLS) на 7 таблицах + автономный отказоустойчивый Demo-режим с локальной памятью."),
        ("Качество кода (100% Green CI):", "Все 14 юнит-тестов FAO-56, 7 интеграционных тестов и 10-шаговый E2E-сценарий фермера успешно проходят в автоматическом GitHub Actions CI."),
        ("Полноценный учёт расхода воды:", "Реализовано 3 метода фиксации (расходомер до/после, подача насоса × время, ручной талон) с редактированием, удалением и расчётом план-факт отклонений."),
        ("План пилота на сезон 2026:", "Подготовка полевых испытаний с 3–5 фермерскими хозяйствами в Туркестанской и Кызылординской областях на площади 200–300 га для доказательства физической экономии воды.")
    ]
    
    for title, desc in items6:
        p_item = tf6.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(8)

    # =========================================================================
    # SLIDE 7: Конкуренты (Таблица)
    # =========================================================================
    slide7 = prs.slides[6]
    format_title(slide7.shapes[0], "7. Конкуренты: Почему клиент выберет AquaAI")
    
    tf7_sub = slide7.shapes[1].text_frame
    tf7_sub.clear()
    p = tf7_sub.paragraphs[0]
    p.text = "Сравнительный анализ решений по управлению орошением в Центральной Азии:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_MUTED
    
    # Table shape 2
    table = slide7.shapes[2].table
    
    # Configure Table Content
    headers = ["Решение", "Цена для фермера", "Ключевое преимущество", "Слабое место / Ограничение"]
    row_data = [
        ("AquaAI", "$1.50 / га в сезон (~750 ₸)", "Работает в браузере без датчиков, казахский язык, расчёт по насосам", "Молодой бренд на стадии запуска пилотных испытаний"),
        ("Зарубежные системы (CropX, FieldNET)", "$3,000–$10,000 за комплект", "Высокая точность при наличии дорогого оборудования", "Недоступно малым фермерам, нет казахского языка, сложный монтаж"),
        ("Традиционный арык («на глаз»)", "0 ₸ за софт (но миллионы за перелив)", "Привычно фермерам, не требует обучения", "Потери до 40% воды, переплаты за насосы, риск гибели урожая")
    ]
    
    # Fill headers
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.text = h_text
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER
        
    # Fill rows
    for r_idx, (col0, col1, col2, col3) in enumerate(row_data):
        row_cells = [col0, col1, col2, col3]
        is_aqua = (r_idx == 0)
        for c_idx, val in enumerate(row_cells):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            if is_aqua:
                cell.fill.fore_color.rgb = COLOR_HIGHLIGHT
            else:
                cell.fill.fore_color.rgb = COLOR_WHITE if r_idx % 2 == 0 else COLOR_LIGHT_BG
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Segoe UI"
            p.font.size = Pt(12)
            p.font.bold = is_aqua
            p.font.color.rgb = COLOR_ACCENT if (is_aqua and c_idx == 0) else COLOR_DARK

    # =========================================================================
    # SLIDE 8: Команда
    # =========================================================================
    slide8 = prs.slides[7]
    format_title(slide8.shapes[0], "8. Команда: Экспертиза для масштабирования")
    
    tf8 = slide8.shapes[1].text_frame
    tf8.clear()
    
    team_members = [
        ("Салават Ахметов — Founder & Lead Full-Stack Architect:", "Опыт в разработке высоконагруженных веб-сервисов (Next.js, TypeScript, PostgreSQL), архитектура масштабируемых систем, интеграция ИИ и облачных API. Полная реализация ядра платформы AquaAI."),
        ("Агрономический эксперт-консультант (AgriTech & Irrigation):", "Экспертиза в почвоведении и прикладной гидрологии орошаемого земледелия Казахстана, практический опыт адаптации математических норм FAO-56 под условия южных областей."),
        ("Инженер по насосному оборудованию и телеметрии:", "Специализация на подборе и эксплуатации глубинных и поверхностных насосных агрегатов, гидравлических расчётах и капельном орошении."),
        ("Менторы и отраслевые советники:", "Поддержка экспертов водного хозяйства и предпринимателей клуба Parasat для выстраивания партнёрств с агрокооперативами.")
    ]
    
    for title, desc in team_members:
        p_item = tf8.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(10)

    # =========================================================================
    # SLIDE 9: Запрос
    # =========================================================================
    slide9 = prs.slides[8]
    format_title(slide9.shapes[0], "9. Запрос: Инвестиции и партнёрства для пилота")
    
    tf9 = slide9.shapes[1].text_frame
    tf9.clear()
    
    items9 = [
        ("Цель на конкурсе IDEA MINERS 2026:", "Привлечение грантового финансирования, менторской поддержки и отраслевых связей для проведения масштабного полевого пилота летом 2026 года."),
        ("Структура расходования средств (100%):", "Распределение бюджета гранта по ключевым статьям развития стартапа:"),
        ("  • 40% — Полевые пилотные испытания:", "Закупка контрольных поверенных расходомеров для оборудования 3–5 пилотных полей в Туркестанской и Кызылординской областях для сбора эталонных данных."),
        ("  • 30% — Развитие продукта и мобильная версия:", "Создание офлайн PWA-клиента для работы поливальщиков без интернета и интеграция спутниковых снимков Sentinel-2."),
        ("  • 30% — Выход на рынок и агро-маркетинг:", "Проведение демонстрационных полевых дней для агрономов и подключение первых 20 коммерческих хозяйств к сезону 2027."),
        ("Партнёры, которые нам нужны:", "Сельхозкооперативы (СПК), крестьянские хозяйства и Ассоциации водопользователей юга Казахстана для совместной верификации эффективности.")
    ]
    
    for idx, (title, desc) in enumerate(items9):
        p_item = tf9.add_paragraph()
        run_t = p_item.add_run()
        run_t.text = title + " "
        run_t.font.name = "Segoe UI"
        run_t.font.size = Pt(14 if idx > 1 else 15)
        run_t.font.bold = True
        run_t.font.color.rgb = COLOR_PRIMARY if idx <= 1 else COLOR_ACCENT
        
        run_d = p_item.add_run()
        run_d.text = desc
        run_d.font.name = "Segoe UI"
        run_d.font.size = Pt(14)
        run_d.font.color.rgb = COLOR_DARK
        p_item.space_after = Pt(6)

    # Save filled presentation
    prs.save(output_pptx)
    print(f"[OK] Presentation saved successfully: {output_pptx}")
    
    # Export to PDF via PowerPoint COM
    print("Exporting presentation to PDF via PowerPoint COM...")
    cmd = f'''powershell -Command "
$ppt = New-Object -ComObject PowerPoint.Application
$file = '{output_pptx}'
$pres = $ppt.Presentations.Open($file, [Microsoft.Office.Core.MsoTriState]::msoTrue, [Microsoft.Office.Core.MsoTriState]::msoFalse, [Microsoft.Office.Core.MsoTriState]::msoFalse)
$out = '{output_pdf}'
if (Test-Path $out) {{ Remove-Item $out -Force }}
$pres.SaveAs($out, 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
Write-Host 'PDF generated:' (Test-Path $out)
"'''
    ret = os.system(cmd)
    if ret == 0 and os.path.exists(output_pdf):
        size_kb = os.path.getsize(output_pdf) / 1024
        print(f"[OK] PDF generated successfully: {output_pdf} ({size_kb:.1f} KB)")
    else:
        print(f"[WARN] Warning: PDF export returned code {ret}")

if __name__ == "__main__":
    create_pitch_deck()
